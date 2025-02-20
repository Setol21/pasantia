import shutil, os, pptx,platform
import pandas as pd

from pathlib import Path

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def create_and_get_ppt_path(retail_name: str, type_name: str, amount_formatted: str, OC_specific_number: str) -> Path:
    ppt_filename = f"{type_name.upper()} {retail_name.upper()} (${amount_formatted.replace(',','.')}) - {OC_specific_number}.pptx"
    
    folder_path = Path.home() / "Desktop" / "PPTs" / f"{retail_name}" / f"{type_name}"
    folder_path.mkdir(parents=True, exist_ok=True)

    return folder_path

def read_and_get() -> None:
    my_data = pd.read_excel("Planilla de Pagos MX.xlsx",sheet_name="Detalle OCs",skiprows=1)

    for i in range(len(my_data["Descripción"])):
        create_ppt(my_data["Descripción"][i],my_data["Retail"][i],int(my_data["Comprometido Retail"][i]),my_data["OC"][i])

def create_ppt(type_name: str, retail_name: str, amount: int , OC_specific_number: str)-> None:
    amount_formatted = f"{amount:,}"

    # Creating the presentation
    prs = pptx.Presentation()

    # Creating the title
    title_slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    title.text = "EVIDENCIAS"

    # Creating the RETAIL box centered
    top = height = Inches(2) 
    width = Inches(5)
    left = (prs.slide_width - width)/2
    retail_box = slide.shapes.add_textbox(left, top, width, height)

    # Textframe inside the box
    tf_retail = retail_box.text_frame

    # Adding the paragraph
    retail = tf_retail.paragraphs[0]
    retail.text = retail_name.upper()
    retail.font.bold = True
    retail.alignment = PP_ALIGN.CENTER

    run = retail.runs[0]
    run.font.size = Pt(72)
    run.font.color.rgb = RGBColor(128,0,64)


    # Add the OC number
    oc_number = tf_retail.add_paragraph()
    oc_number.text = OC_specific_number
    oc_number.alignment = PP_ALIGN.CENTER
    oc_number.font.size = Pt(48)

    # Add the Period
    period = tf_retail.add_paragraph()
    period.text = "Oct 1st to Dec 31st 2024"
    period.alignment = PP_ALIGN.CENTER
    period.font.size = Pt(28)

    prs.save(create_and_get_ppt_path(retail_name,type_name,amount_formatted,OC_specific_number) / f"{type_name.upper()} {type_name.upper()} (${amount_formatted.replace(',','.')}) - {OC_specific_number}.pptx")
    return


def main():
    read_and_get()

if __name__ == "__main__":
    main()